"""
main.py — Sistema de Eye Tracking y Análisis de Sentimientos ECOTEC 2026
Cambios:
  - Formulario completo (código ECOTEC, facultad dinámica, carrera, semestre, gafas, lentes)
  - Dos paneles iguales (comercial + cámara) en columna izquierda
  - Formulario en columna derecha
  - Guardado automático al terminar el video
  - Integración con export_utils.py
"""

import sys
import os
import cv2
import numpy as np
from datetime import datetime

from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QLabel, QLineEdit, QProgressBar, QFrame, QComboBox,
    QDialog, QTableWidget, QTableWidgetItem, QHeaderView, QSpinBox,
    QMessageBox, QCheckBox, QSizePolicy
)
from PyQt6.QtCore import QTimer, Qt, QThread, pyqtSignal
from PyQt6.QtGui  import QImage, QPixmap, QFont

# ── Colores ────────────────────────────────────────────────────────────────────
ROWS = 4
COLS = 4
AZUL_OSCURO  = "#002060"
AZUL_MEDIO   = "#003087"
AZUL_CLARO   = "#0057B8"
VERDE        = "#00A651"
VERDE_OSCURO = "#007A3D"
BLANCO       = "#FFFFFF"
GRIS_CLARO   = "#F0F4F8"
GRIS_MEDIO   = "#D0D8E4"
TEXTO        = "#1A1A2E"
ROJO         = "#C0392B"

def _leer_ultimo_id():
    """Lee el último ID del consolidado y retorna el siguiente número.
    Si no existe el consolidado, retorna 1 (P001)."""
    cons_path = 'data/reporte_consolidado.xlsx'
    try:
        if os.path.exists(cons_path):
            import openpyxl as _opx
            wb = _opx.load_workbook(cons_path, read_only=True)
            if 'Resumen' in wb.sheetnames:
                ws = wb['Resumen']
                ultimo = 0
                for row in ws.iter_rows(min_row=3, values_only=True):
                    val = row[1] if row and len(row) > 1 else None  # columna B = ID
                    if val and str(val).startswith('P'):
                        try:
                            num = int(str(val).replace('P', ''))
                            if num > ultimo:
                                ultimo = num
                        except ValueError:
                            pass
                wb.close()
                if ultimo > 0:
                    print(f"[Contador] Último ID en consolidado: P{ultimo:03d} → siguiente: P{ultimo+1:03d}")
                    return ultimo + 1
    except Exception as e:
        print(f"[Contador] Error leyendo consolidado: {e}")
    print("[Contador] Consolidado no encontrado o vacío → iniciando desde P001")
    return 1

participant_counter = [_leer_ultimo_id()]

# ── Facultades y carreras ECOTEC 2025 ──────────────────────────────────────────
FACULTADES_CARRERAS = {
    "Cs. de la Salud y Desarrollo Humano": [
        "Enfermería", "Fisioterapia", "Medicina",
        "Nutrición y Dietética", "Psicología Clínica",
    ],
    "Cs. Económicas y Empresariales": [
        "Administración de Empresas", "Contabilidad y Auditoría", "Economía",
        "Finanzas", "Logística y Transporte", "Gestión del Talento Humano",
        "Negocios Digitales", "Negocios Internacionales",
    ],
    "Comunicación, Humanidades y Creatividad": [
        "Comunicación", "Marketing", "Multimedia y Producción Audiovisual",
        "Psicología", "Periodismo",
    ],
    "Derecho y Gobernabilidad": ["Criminalística", "Derecho"],
    "Estudios Globales y Hospitalidad": ["Relaciones Internacionales", "Turismo y Hotelería"],
    "Ingenierías, Arquitectura y Cs. de la Naturaleza": [
        "Arquitectura", "Agronomía", "Ingeniería Civil", "Ingeniería Industrial",
        "Mecatrónica", "Medicina Veterinaria", "Sistemas Inteligentes",
    ],
    "ECOTEC Online": [
        "Ciencias de la Educación", "Administración de Empresas", "Comercio",
        "Comunicación", "Contabilidad y Auditoría", "Derecho", "Economía",
        "Finanzas", "Gestión del Talento Humano", "Mercadotecnia",
        "Negocios Internacionales", "Psicología", "Sistemas Inteligentes",
        "Trabajo Social", "Turismo",
    ],
}
SEMESTRES = ["Primero","Segundo","Tercero","Cuarto","Quinto",
             "Sexto","Séptimo","Octavo","Noveno","Décimo"]

LOGO_PATH = os.path.join(os.path.dirname(__file__), "assets", "ecotec_logo.png")


def _nombre_safe(nombre):
    """Elimina tildes, ñ y caracteres especiales para nombres de archivo."""
    import unicodedata
    if not nombre:
        return 'SinNombre'
    nfkd     = unicodedata.normalize('NFKD', nombre)
    ascii_str = nfkd.encode('ASCII', 'ignore').decode('ASCII')
    return ascii_str.replace(' ', '_').strip('_') or 'SinNombre'


# ─────────────────────────────────────────────
#  VENTANA DE COMERCIAL EN PANTALLA COMPLETA
# ─────────────────────────────────────────────
class FullscreenVideoWindow(QWidget):
    def __init__(self, screen_index=0, parent=None):
        super().__init__(parent, Qt.WindowType.Window)
        self.setWindowFlags(
            Qt.WindowType.Window |
            Qt.WindowType.FramelessWindowHint |
            Qt.WindowType.WindowStaysOnTopHint
        )
        self.setStyleSheet("background-color: black;")
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        self.video_label = QLabel()
        self.video_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.video_label.setStyleSheet("background-color: black;")
        layout.addWidget(self.video_label)
        self.debug_label = QLabel("FPS: --  |  Sector: --  |  --", self)
        self.debug_label.setStyleSheet(
            "background-color: rgba(0,0,0,180); color: #00FF88; "
            "font-size: 16px; font-weight: bold; padding: 6px 12px; border-radius: 6px;")
        self.debug_label.setFixedHeight(36)
        self.debug_label.adjustSize()
        self.debug_label.move(10, 10)
        self.debug_label.raise_()
        self.debug_label.show()
        screens = QApplication.screens()
        screen  = screens[screen_index] if screen_index < len(screens) else screens[0]
        self.setGeometry(screen.geometry())
        self.showFullScreen()

    def update_frame(self, qt_image):
        lw, lh = self.video_label.width(), self.video_label.height()
        if lw <= 0 or lh <= 0:
            return
        self.video_label.setPixmap(QPixmap.fromImage(
            qt_image.scaled(lw, lh,
                            Qt.AspectRatioMode.KeepAspectRatio,
                            Qt.TransformationMode.FastTransformation)))

    def update_debug_info(self, sector, yaw, pitch, emotion, fps=0.0):
        self.debug_label.setText(
            f"FPS: {fps:.0f}  |  Sector: {sector}  |  "
            f"Yaw: {yaw:.1f}  Pitch: {pitch:.1f}  |  {emotion}")
        self.debug_label.adjustSize()

    def get_render_size(self):
        return self.video_label.width(), self.video_label.height()

    def close_window(self):
        self.hide()
        self.deleteLater()


# ─────────────────────────────────────────────
#  DIÁLOGO DE DISTANCIA
# ─────────────────────────────────────────────
class DistanceDialog(QDialog):
    FOCAL_PX = 1101.5
    FACE_CM  = 14.0

    def __init__(self, controller, parent=None):
        super().__init__(parent)
        self.controller       = controller
        self.setWindowTitle("Configurar calibración")
        self.setFixedSize(440, 340)
        self.setStyleSheet(f"background-color: {GRIS_CLARO};")
        self.distance_cm      = 90
        self.camera_height_cm = 110
        self.eye_height_cm    = 110
        self._build_ui()

    def _build_ui(self):
        layout = QVBoxLayout(self)
        layout.setSpacing(12)
        layout.setContentsMargins(24, 20, 24, 20)
        title = QLabel("📏  Configuración antes de calibrar")
        title.setFont(QFont("Arial", 12, QFont.Weight.Bold))
        title.setStyleSheet(f"color: {AZUL_OSCURO};")
        layout.addWidget(title)
        sep = QFrame()
        sep.setFrameShape(QFrame.Shape.HLine)
        sep.setStyleSheet(f"color: {GRIS_MEDIO};")
        layout.addWidget(sep)
        dist_row = QHBoxLayout()
        lbl_d = QLabel("Distancia al participante (automática):")
        lbl_d.setStyleSheet(f"color: {TEXTO}; font-size: 10px;")
        lbl_d.setFixedWidth(240)
        dist_row.addWidget(lbl_d)
        self.lbl_dist = QLabel("Detectando...")
        self.lbl_dist.setStyleSheet(f"color: {VERDE}; font-size: 12px; font-weight: bold;")
        dist_row.addWidget(self.lbl_dist)
        layout.addLayout(dist_row)
        for lbl_text, attr, lo, hi, val, suffix in [
            ("Altura de la cámara (desde el piso):", "spin_cam", 50, 300, 110, " cm"),
            ("Altura de ojos del participante sentado:", "spin_eye", 80, 160, 110, " cm"),
        ]:
            row = QHBoxLayout()
            lbl = QLabel(lbl_text)
            lbl.setStyleSheet(f"color: {TEXTO}; font-size: 10px;")
            lbl.setFixedWidth(240)
            row.addWidget(lbl)
            spin = QSpinBox()
            spin.setRange(lo, hi); spin.setValue(val); spin.setSuffix(suffix)
            spin.setFixedWidth(100)
            spin.setStyleSheet(
                f"padding: 4px; border: 1px solid {GRIS_MEDIO}; border-radius: 4px; font-size: 11px;")
            setattr(self, attr, spin)
            row.addWidget(spin)
            layout.addLayout(row)
        self.lbl_info = QLabel("")
        self.lbl_info.setStyleSheet(f"color: {AZUL_CLARO}; font-size: 9px;")
        self.lbl_info.setWordWrap(True)
        layout.addWidget(self.lbl_info)
        self.spin_cam.valueChanged.connect(self._update_info)
        self.spin_eye.valueChanged.connect(self._update_info)
        btn_ok = QPushButton("✓  CONFIRMAR Y CALIBRAR")
        btn_ok.setFixedHeight(40)
        btn_ok.setStyleSheet(f"""
            QPushButton {{ background-color:{AZUL_CLARO}; color:{BLANCO};
                font-size:11px; font-weight:bold; border-radius:8px; }}
            QPushButton:hover {{ background-color:{AZUL_MEDIO}; }}""")
        btn_ok.clicked.connect(self._confirm)
        layout.addWidget(btn_ok)
        QTimer.singleShot(300, self._detect_distance)

    def _detect_distance(self):
        try:
            cap     = self.controller.cap
            cascade = self.controller.face_cascade
            if cap is None or not cap.isOpened():
                self.lbl_dist.setText("No disponible")
                self.lbl_dist.setStyleSheet(f"color:{ROJO}; font-size:11px;")
                self._update_info(); return
            face_widths = []
            for _ in range(40):
                ret, frame = cap.read()
                if not ret: continue
                small = cv2.resize(frame, (640, 400))
                gray  = cv2.cvtColor(small, cv2.COLOR_BGR2GRAY)
                faces = cascade.detectMultiScale(gray, 1.1, 5, minSize=(40,40))
                if len(faces) > 0:
                    fx, fy, fw, fh = max(faces, key=lambda f: f[2]*f[3])
                    face_widths.append(fw * (1920/640))
            if len(face_widths) >= 5:
                median_px = float(np.median(face_widths))
                self.distance_cm = max(20, min(500, int((self.FOCAL_PX*self.FACE_CM)/median_px)))
                self.lbl_dist.setText(f"{self.distance_cm} cm")
                self.lbl_dist.setStyleSheet(f"color:{VERDE}; font-size:12px; font-weight:bold;")
            else:
                self.lbl_dist.setText("Rostro no detectado")
                self.lbl_dist.setStyleSheet(f"color:{ROJO}; font-size:11px;")
                self.distance_cm = 90
        except Exception as e:
            print(f"Error detectando distancia: {e}")
            self.distance_cm = 90
            self.lbl_dist.setText("Error")
        self._update_info()

    def _update_info(self):
        import math
        dist, cam_h, eye_h = self.distance_cm, self.spin_cam.value(), self.spin_eye.value()
        po = math.degrees(math.atan2(cam_h-eye_h, dist))
        yt = max(25, min(45, int(30+(dist-60)/20)))
        pt = max(20, min(40, int(25+(dist-60)/15)))
        self.lbl_info.setText(
            f"Distancia: {dist}cm  |  Pitch offset: {po:.1f}°  |  Umbral yaw: ±{yt}°  pitch: ±{pt}°")

    def _confirm(self):
        self.camera_height_cm = self.spin_cam.value()
        self.eye_height_cm    = self.spin_eye.value()
        self.accept()

    def get_geometry_params(self):
        import math
        dist, cam_h, eye_h = self.distance_cm, self.camera_height_cm, self.eye_height_cm
        po = math.degrees(math.atan2(cam_h-eye_h, dist))
        return {
            'pitch_offset':     po,
            'yaw_threshold':    max(25, min(45, int(30+(dist-60)/20))),
            'pitch_threshold':  max(20, min(40, int(25+(dist-60)/15))),
            'distance_cm':      dist,
        }


# ─────────────────────────────────────────────
#  THREADS
# ─────────────────────────────────────────────
class LoadingThread(QThread):
    progress = pyqtSignal(str)
    finished = pyqtSignal(object)

    def run(self):
        self.progress.emit("Cargando modelos de IA...")
        from controllers.session_controller import SessionController
        controller = SessionController()
        self.progress.emit("Inicializando L2CS-Net y HSEmotions...")
        controller.initialize()
        controller.pre_extracted_audio = None
        controller.calibration_matrix  = None
        self.progress.emit("Sistema listo.")
        self.finished.emit(controller)


class AudioThread(QThread):
    finished = pyqtSignal(str)

    def __init__(self, video_path):
        super().__init__()
        self.video_path = video_path

    def run(self):
        try:
            from moviepy import VideoFileClip
            audio_path = os.path.splitext(self.video_path)[0] + '.mp3'
            if not os.path.exists(audio_path):
                clip = VideoFileClip(self.video_path)
                clip.audio.write_audiofile(audio_path)
                clip.close()
            self.finished.emit(audio_path)
        except Exception as e:
            print(f"Error extrayendo audio: {e}")
            self.finished.emit("")


# ─────────────────────────────────────────────
#  DIÁLOGO DE RESULTADOS (solo visualización)
# ─────────────────────────────────────────────
class ResultsDialog(QDialog):
    def __init__(self, summary, heatmap, sector_log,
                 participant_name="", saved_files=None, parent=None, aoi_log=None):
        super().__init__(parent)
        self.setWindowTitle("Resultados de la Sesión")
        self.setMinimumSize(960, 680)
        self.setStyleSheet(f"background-color: {GRIS_CLARO};")
        self._summary = summary
        saved_files   = saved_files or {}
        aoi_log       = aoi_log or []

        layout = QVBoxLayout(self)
        layout.setSpacing(10)
        layout.setContentsMargins(18, 18, 18, 18)

        # ── Encabezado ────────────────────────────────────────────
        title_lbl = QLabel("📊  ANÁLISIS COMPLETO DE ENGAGEMENT PUBLICITARIO")
        title_lbl.setFont(QFont("Arial", 13, QFont.Weight.Bold))
        title_lbl.setStyleSheet(f"color: {AZUL_OSCURO};")
        layout.addWidget(title_lbl)

        if participant_name:
            name_lbl = QLabel(f"👤  {participant_name}  —  {summary.get('participante','—')}")
            name_lbl.setFont(QFont("Arial", 11))
            name_lbl.setStyleSheet(f"color: {AZUL_CLARO};")
            layout.addWidget(name_lbl)

        sep = QFrame(); sep.setFrameShape(QFrame.Shape.HLine)
        sep.setStyleSheet(f"color: {AZUL_CLARO};"); layout.addWidget(sep)

        top = QHBoxLayout()

        # ── Columna izquierda: métricas ───────────────────────────
        left = QVBoxLayout(); left.setSpacing(5)
        emotion_icons = {'happy':'😊','sad':'😢','angry':'😠',
                         'fear':'😨','surprise':'😲','disgust':'🤢','neutral':'😐'}
        metricas_title = QLabel("MÉTRICAS GENERALES")
        metricas_title.setFont(QFont("Arial",10,QFont.Weight.Bold))
        metricas_title.setStyleSheet(f"color: {AZUL_MEDIO};")
        left.addWidget(metricas_title)

        # Primera fijación al AOI: mismo cálculo que en el PDF/Excel — el segundo
        # más bajo entre las 3 AOIs (1 estática + 2 dinámicas), no el genérico
        # de "mirar la pantalla" (que casi siempre da ~0.00s y no dice nada útil).
        _primera_fij_aoi = f"{min(e['segundo'] for e in aoi_log)}s" if aoi_log else "—"

        segs = summary.get('engagement_segmentos', {})
        for key, val in [
            ("📈 Engagement promedio",     f"{summary.get('engagement_promedio',0)*100:.1f}%"),
            ("📈 Engagement máximo",       f"{summary.get('engagement_maximo',0)*100:.1f}%"),
            ("😊 Emoción dominante",       summary.get('emocion_dominante','—')),
            ("👁 Tiempo mirando pantalla", f"{summary.get('tiempo_mirando',0):.1f}%"),
            ("⚡ Primera fijación al AOI", _primera_fij_aoi),
            ("😵 Distracciones",           str(summary.get('num_distracciones',0))),
            ("💚 Valencia emocional",      f"{summary.get('valencia_emocional',0):+.1f}%"),
            ("🔄 Variabilidad emocional",  f"{summary.get('variabilidad_emocional',0):.1f}/min"),
            ("Seg. 1 (0–15s)",            f"{segs.get('segmento_1',0)*100:.1f}%"),
            ("Seg. 2 (15–30s)",           f"{segs.get('segmento_2',0)*100:.1f}%"),
            ("Seg. 3 (30–45s)",           f"{segs.get('segmento_3',0)*100:.1f}%"),
            ("Seg. 4 (45–60s)",           f"{segs.get('segmento_4',0)*100:.1f}%"),
        ]:
            row = QWidget()
            row.setStyleSheet(f"background-color:{BLANCO}; border-radius:5px; border:1px solid {GRIS_MEDIO};")
            rl = QHBoxLayout(row); rl.setContentsMargins(8,4,8,4)
            k = QLabel(key+":"); k.setFont(QFont("Arial",9,QFont.Weight.Bold))
            k.setStyleSheet(f"color:{AZUL_MEDIO}; border:none;"); k.setFixedWidth(190); rl.addWidget(k)
            v = QLabel(val); v.setFont(QFont("Arial",9))
            v.setStyleSheet(f"color:{TEXTO}; border:none;"); rl.addWidget(v)
            left.addWidget(row)

        eng_lbl = QLabel("Índice de Engagement:")
        eng_lbl.setStyleSheet(f"color:{AZUL_MEDIO}; font-weight:bold; font-size:9px;")
        left.addWidget(eng_lbl)
        eng_bar = QProgressBar()
        eng_bar.setRange(0,100); eng_bar.setValue(int(summary.get('engagement_promedio',0)*100))
        eng_bar.setFixedHeight(18)
        eng_bar.setStyleSheet(f"""QProgressBar {{ border:1px solid {GRIS_MEDIO};
            border-radius:5px; background:{GRIS_CLARO}; text-align:center; font-weight:bold; }}
            QProgressBar::chunk {{ background:{VERDE}; border-radius:5px; }}""")
        left.addWidget(eng_bar)

        if sector_log:
            import pandas as pd
            df_s = pd.DataFrame(sector_log)
            st = QLabel("TOP SECTORES")
            st.setFont(QFont("Arial",10,QFont.Weight.Bold))
            st.setStyleSheet(f"color:{AZUL_MEDIO};"); left.addWidget(st)
            for sec, cnt in df_s['sector'].value_counts().head(3).items():
                pct = cnt/len(df_s)*100
                row = QWidget()
                row.setStyleSheet(f"background-color:{BLANCO}; border-radius:5px; border:1px solid {GRIS_MEDIO};")
                rl = QHBoxLayout(row); rl.setContentsMargins(8,4,8,4)
                k = QLabel(f"🎯 Sector {sec}:"); k.setStyleSheet(f"color:{AZUL_MEDIO}; border:none; font-size:9px;"); k.setFixedWidth(80); rl.addWidget(k)
                v = QLabel(f"{cnt}s ({pct:.1f}%)"); v.setStyleSheet(f"color:{TEXTO}; border:none; font-size:9px;"); rl.addWidget(v)
                left.addWidget(row)

        left.addStretch()
        top.addLayout(left)

        # ── Columna derecha: mapas visuales ──────────────────────
        right = QVBoxLayout(); right.setSpacing(5)
        hm_title = QLabel("🌡  MAPA DE CALOR")
        hm_title.setFont(QFont("Arial",10,QFont.Weight.Bold))
        hm_title.setStyleSheet(f"color:{AZUL_OSCURO};")
        hm_title.setAlignment(Qt.AlignmentFlag.AlignCenter); right.addWidget(hm_title)
        hm_lbl = QLabel(); hm_lbl.setFixedSize(400,225)
        hm_lbl.setStyleSheet(f"border:2px solid {AZUL_MEDIO}; border-radius:6px;")
        if heatmap is not None and heatmap.max() > 0:
            hn = cv2.normalize(heatmap,None,0,255,cv2.NORM_MINMAX)
            hc = cv2.applyColorMap(hn.astype(np.uint8), cv2.COLORMAP_JET)
            hr = cv2.resize(hc,(400,225)); hrg = cv2.cvtColor(hr,cv2.COLOR_BGR2RGB)
            h, w, ch = hrg.shape
            hm_lbl.setPixmap(QPixmap.fromImage(QImage(hrg.data,w,h,ch*w,QImage.Format.Format_RGB888)))
        else:
            hm_lbl.setText("Sin datos"); hm_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        right.addWidget(hm_lbl)

        sm_title = QLabel("🗺  MAPA DE SECTORES")
        sm_title.setFont(QFont("Arial",10,QFont.Weight.Bold))
        sm_title.setStyleSheet(f"color:{AZUL_OSCURO};")
        sm_title.setAlignment(Qt.AlignmentFlag.AlignCenter); right.addWidget(sm_title)
        sm_lbl = QLabel(); sm_lbl.setFixedSize(400,160)
        sm_lbl.setStyleSheet(f"border:2px solid {AZUL_MEDIO}; border-radius:6px;")
        if sector_log:
            import pandas as pd
            df_sm = pd.DataFrame(sector_log)
            simg  = np.zeros((160,400,3),dtype=np.uint8); simg[:] = (30,30,30)
            cw,chc = 400//COLS,160//ROWS
            scm = df_sm['sector'].value_counts(); mx = scm.max() if len(scm)>0 else 1
            for r in range(ROWS):
                for c in range(COLS):
                    sn=r*COLS+c+1; cnt=scm.get(sn,0)
                    x1,y1=c*cw,r*chc; x2,y2=x1+cw-2,y1+chc-2
                    cv2.rectangle(simg,(x1,y1),(x2,y2),(0,int(cnt/mx*255),255-int(cnt/mx*255)),-1)
                    cv2.rectangle(simg,(x1,y1),(x2,y2),(255,255,255),1)
                    cv2.putText(simg,str(sn),(x1+5,y1+20),cv2.FONT_HERSHEY_SIMPLEX,0.5,(255,255,255),1)
                    if cnt>0: cv2.putText(simg,f"{cnt}s",(x1+5,y1+38),cv2.FONT_HERSHEY_SIMPLEX,0.38,(255,255,0),1)
            sg = cv2.cvtColor(simg,cv2.COLOR_BGR2RGB); h,w,ch = sg.shape
            sm_lbl.setPixmap(QPixmap.fromImage(QImage(sg.data,w,h,ch*w,QImage.Format.Format_RGB888)))
        else:
            sm_lbl.setText("Sin datos"); sm_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        right.addWidget(sm_lbl)
        right.addStretch()
        top.addLayout(right)
        layout.addLayout(top)

        # ── Archivos guardados ────────────────────────────────────
        if saved_files:
            saved_lbl = QLabel("✅  ARCHIVOS GUARDADOS AUTOMÁTICAMENTE")
            saved_lbl.setFont(QFont("Arial",10,QFont.Weight.Bold))
            saved_lbl.setStyleSheet(f"color:{VERDE};")
            layout.addWidget(saved_lbl)
            for tipo, ruta in saved_files.items():
                r = QWidget()
                r.setStyleSheet(f"background-color:{BLANCO}; border-radius:5px; border:1px solid {GRIS_MEDIO};")
                rl = QHBoxLayout(r); rl.setContentsMargins(10,4,10,4)
                k = QLabel(tipo+":"); k.setFont(QFont("Arial",9,QFont.Weight.Bold))
                k.setStyleSheet(f"color:{AZUL_MEDIO}; border:none;"); k.setFixedWidth(120); rl.addWidget(k)
                v = QLabel(ruta); v.setStyleSheet(f"color:{TEXTO}; border:none; font-size:9px;"); rl.addWidget(v)
                layout.addWidget(r)

        # ── Registro ──────────────────────────────────────────────
        if sector_log:
            tbl_t = QLabel("⏱  REGISTRO SEGUNDO A SEGUNDO")
            tbl_t.setFont(QFont("Arial",10,QFont.Weight.Bold))
            tbl_t.setStyleSheet(f"color:{AZUL_OSCURO};"); layout.addWidget(tbl_t)
            tabla = QTableWidget(); tabla.setMaximumHeight(130); tabla.setColumnCount(3)
            tabla.setHorizontalHeaderLabels(["Segundo","Sector mirado","Emoción"])
            tabla.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
            tabla.setStyleSheet(f"background-color:{BLANCO}; border:1px solid {GRIS_MEDIO};")
            tabla.setRowCount(len(sector_log))
            for i,entry in enumerate(sector_log):
                tabla.setItem(i,0,QTableWidgetItem(f"{entry['tiempo_seg']}s"))
                tabla.setItem(i,1,QTableWidgetItem(f"Sector {entry['sector']}"))
                icon = emotion_icons.get(entry['emocion'],'❓')
                tabla.setItem(i,2,QTableWidgetItem(f"{icon} {entry['emocion']}"))
            layout.addWidget(tabla)

        # ── Botón cerrar ──────────────────────────────────────────
        btn_close = QPushButton("CERRAR")
        btn_close.setFixedHeight(40)
        btn_close.setStyleSheet(f"""QPushButton {{ background-color:{AZUL_CLARO}; color:{BLANCO};
            font-size:12px; font-weight:bold; border-radius:6px; }}
            QPushButton:hover {{ background-color:{AZUL_MEDIO}; }}""")
        btn_close.clicked.connect(self.close)
        layout.addWidget(btn_close)


# ─────────────────────────────────────────────
#  VISTA PRINCIPAL
# ─────────────────────────────────────────────
class MainView(QMainWindow):
    def __init__(self, controller):
        super().__init__()
        self.controller            = controller
        self.camera_thread         = None
        self.calibration_widget    = None
        self.fullscreen_video      = None
        self.heatmap               = np.zeros((400, 640), dtype=np.float32)
        self._heatmap_rw           = 640
        self._heatmap_rh           = 400
        self.video_player_obj      = None
        self.audio_thread          = None
        self.session_active        = False
        self._participant_name     = ""
        self._participant_codigo   = ""
        self._participant_edad     = ""
        self._participant_sexo     = ""
        self._participant_facultad = ""
        self._participant_carrera  = ""
        self._participant_semestre = ""
        self._participant_gafas    = False
        self._participant_lentes   = False
        self._participant_distancia= 90
        self._participant_photo    = None

        self.setWindowTitle("Sistema de Eye Tracking y Análisis de Sentimientos — ECOTEC")
        self.showMaximized()
        self.setStyleSheet(f"background-color: {GRIS_CLARO};")
        self._build_ui()

    def _target_screen_index(self):
        return 1 if len(QApplication.screens()) > 1 else 0

    def start_camera_thread(self):
        from controllers.session_controller import CameraThread
        self.camera_thread = CameraThread(self.controller)
        self.camera_thread.frame_ready.connect(self.on_camera_frame)
        self.camera_thread.start()

    def _capture_photo(self, participant_id, nombre):
        """Captura foto con barra azul usando PIL (soporta tildes y ñ)."""
        try:
            from PIL import Image, ImageDraw, ImageFont
            os.makedirs('data', exist_ok=True)
            ret, frame = self.controller.cap.read()
            if not ret:
                return None

            frame_flip = cv2.flip(frame, 1)
            h, w      = frame_flip.shape[:2]

            # Convertir BGR → RGB para PIL
            frame_rgb = cv2.cvtColor(frame_flip, cv2.COLOR_BGR2RGB)
            pil_img   = Image.fromarray(frame_rgb)

            # Barra azul con texto usando PIL (soporta tildes y ñ)
            bar  = Image.new('RGB', (w, 52), (0, 32, 96))
            draw = ImageDraw.Draw(bar)

            font = None
            for fp in [r'C:\Windows\Fonts\arial.ttf',
                       r'C:\Windows\Fonts\calibri.ttf',
                       r'C:\Windows\Fonts\segoeui.ttf']:
                try:
                    font = ImageFont.truetype(fp, 26)
                    break
                except Exception:
                    pass
            if font is None:
                font = ImageFont.load_default()

            draw.text((16, 13), f"{nombre}  |  {participant_id}",
                      fill=(255, 255, 255), font=font)

            # Combinar frame + barra
            combined = Image.new('RGB', (w, h + 52))
            combined.paste(pil_img, (0, 0))
            combined.paste(bar,     (0, h))

            safe_name = _nombre_safe(nombre)
            path = f"data/foto_{participant_id}_{safe_name}.jpg"
            combined.save(path, 'JPEG', quality=95)
            return path

        except Exception as e:
            print(f"Error capturando foto: {e}")
        return None

    # ── panel helper ─────────────────────────────────────────────
    def _panel_header(self, title, badge_text="", badge_color=VERDE):
        w = QWidget()
        w.setStyleSheet(f"background-color: {AZUL_MEDIO}; border-radius: 6px 6px 0 0;")
        l = QHBoxLayout(w); l.setContentsMargins(10, 5, 10, 5)
        t = QLabel(title); t.setFont(QFont("Arial",10,QFont.Weight.Bold))
        t.setStyleSheet(f"color: {BLANCO};"); l.addWidget(t)
        if badge_text:
            b = QLabel(badge_text); b.setFont(QFont("Arial",10))
            b.setStyleSheet(f"color: {badge_color}; border:none;"); l.addWidget(b)
        return w

    def _metrics_bar(self, attrs_labels):
        """
        attrs_labels: [(attr_name, display_text), ...]
        Devuelve el widget y guarda las labels en self.attr_name
        """
        w = QWidget()
        w.setStyleSheet(f"background-color: {BLANCO}; border-top: 1.5px solid {AZUL_MEDIO};")
        l = QHBoxLayout(w); l.setContentsMargins(0,0,0,0); l.setSpacing(0)
        for attr, text in attrs_labels:
            cell = QWidget()
            cell.setStyleSheet(f"border-right: 0.5px solid {GRIS_MEDIO};")
            cl = QVBoxLayout(cell); cl.setContentsMargins(4,4,4,4); cl.setSpacing(1)
            lbl_name = QLabel(text.upper())
            lbl_name.setStyleSheet(f"font-size:8px; color:#888; border:none;")
            lbl_name.setAlignment(Qt.AlignmentFlag.AlignCenter)
            lbl_val  = QLabel("—")
            lbl_val.setFont(QFont("Arial",12,QFont.Weight.Medium))
            lbl_val.setStyleSheet(f"color:{TEXTO}; border:none;")
            lbl_val.setAlignment(Qt.AlignmentFlag.AlignCenter)
            setattr(self, attr, lbl_val)
            cl.addWidget(lbl_name); cl.addWidget(lbl_val)
            l.addWidget(cell, 1)
        return w

    def _eng_strip(self, attr_bar, attr_lbl_title, title="Índice de engagement en tiempo real"):
        w = QWidget()
        w.setStyleSheet(f"background-color:{BLANCO}; border-top:0.5px solid {GRIS_MEDIO};")
        l = QVBoxLayout(w); l.setContentsMargins(10,5,10,5); l.setSpacing(3)
        lbl = QLabel(title); lbl.setStyleSheet(f"color:{AZUL_MEDIO}; font-weight:bold; font-size:9px;")
        setattr(self, attr_lbl_title, lbl); l.addWidget(lbl)
        bar = QProgressBar(); bar.setRange(0,100); bar.setValue(0); bar.setFixedHeight(14)
        bar.setStyleSheet(f"""QProgressBar {{ border:1px solid {GRIS_MEDIO}; border-radius:4px;
            background:{GRIS_CLARO}; text-align:center; font-size:9px; }}
            QProgressBar::chunk {{ background:{VERDE}; border-radius:4px; }}""")
        setattr(self, attr_bar, bar); l.addWidget(bar)
        return w

    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        main_layout = QVBoxLayout(central)
        main_layout.setSpacing(0)
        main_layout.setContentsMargins(0, 0, 0, 0)

        # ── TOPBAR ────────────────────────────────────────────────
        header = QWidget()
        header.setFixedHeight(52)
        header.setStyleSheet(f"background-color:{AZUL_OSCURO}; border-bottom:3px solid {VERDE};")
        hl = QHBoxLayout(header); hl.setContentsMargins(14, 0, 14, 0); hl.setSpacing(10)

        # Logo ECOTEC
        logo_lbl = QLabel()
        logo_lbl.setFixedSize(38, 38)
        if os.path.exists(LOGO_PATH):
            pix = QPixmap(LOGO_PATH).scaled(38, 38,
                Qt.AspectRatioMode.KeepAspectRatio,
                Qt.TransformationMode.SmoothTransformation)
            logo_lbl.setPixmap(pix)
            logo_lbl.setStyleSheet("background-color:white; border-radius:5px; padding:2px;")
        else:
            logo_lbl.setText("ECO")
            logo_lbl.setStyleSheet("background-color:white; border-radius:5px; color:#002060; font-weight:bold; font-size:9px;")
            logo_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        hl.addWidget(logo_lbl)

        title_col = QVBoxLayout(); title_col.setSpacing(0)
        ht = QLabel("Universidad ECOTEC  ·  Sistema de Medición de Engagement Publicitario")
        ht.setFont(QFont("Arial",11,QFont.Weight.Bold)); ht.setStyleSheet(f"color:{BLANCO};")
        title_col.addWidget(ht)
        hs = QLabel("Ingeniería en Sistemas Inteligentes  ·  Campus Samborondón 2026")
        hs.setFont(QFont("Arial",9)); hs.setStyleSheet(f"color:{GRIS_MEDIO};")
        title_col.addWidget(hs)
        hl.addLayout(title_col); hl.addStretch()
        self.lbl_id_header = QLabel("")
        self.lbl_id_header.setFont(QFont("Arial",10))
        self.lbl_id_header.setStyleSheet(f"color:{VERDE};"); hl.addWidget(self.lbl_id_header)
        main_layout.addWidget(header)

        # ── CONTENT ───────────────────────────────────────────────
        content = QWidget()
        cl = QHBoxLayout(content); cl.setSpacing(8); cl.setContentsMargins(8,8,8,8)

        # ── COLUMNA IZQUIERDA: dos paneles iguales ────────────────
        left_widget = QWidget()
        left_layout = QVBoxLayout(left_widget)
        left_layout.setSpacing(8); left_layout.setContentsMargins(0,0,0,0)

        # Panel 1: Comercial
        panel_com = QFrame()
        panel_com.setStyleSheet(f"border: 1.5px solid {AZUL_MEDIO}; border-radius: 8px;")
        pcl = QVBoxLayout(panel_com); pcl.setSpacing(0); pcl.setContentsMargins(0,0,0,0)
        pcl.addWidget(self._panel_header("🎬  Comercial — Vista del participante", "▶ 00:00", VERDE))
        self.comercial_label = QLabel()
        self.comercial_label.setStyleSheet(f"background-color:#111; border:none;")
        self.comercial_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.comercial_label.setText("▶  Calibra y luego inicia sesión para reproducir el comercial")
        self.comercial_label.setFont(QFont("Arial",10))
        self.comercial_label.setStyleSheet(f"background-color:#111; color:{GRIS_MEDIO}; border:none;")
        pcl.addWidget(self.comercial_label, 1)
        pcl.addWidget(self._metrics_bar([
            ('mc_emocion','Emoción'),('mc_yaw','Yaw'),('mc_pitch','Pitch'),
            ('mc_pantalla','Pantalla'),('mc_sector','Sector'),('mc_fps','FPS'),
        ]))
        pcl.addWidget(self._eng_strip('engagement_bar','_eng_lbl1'))
        left_layout.addWidget(panel_com, 1)

        # Panel 2: Cámara
        panel_cam = QFrame()
        panel_cam.setStyleSheet(f"border: 1.5px solid {AZUL_MEDIO}; border-radius: 8px;")
        pca = QVBoxLayout(panel_cam); pca.setSpacing(0); pca.setContentsMargins(0,0,0,0)
        pca.addWidget(self._panel_header("📷  Cámara Brio 4K — Eye Tracking en tiempo real", "● REC", ROJO))
        self.camera_label = QLabel()
        self.camera_label.setStyleSheet("background-color:black; border:none;")
        self.camera_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        pca.addWidget(self.camera_label, 1)
        pca.addWidget(self._metrics_bar([
            ('mc_dist','Distancia'),('mc_detec','Detección'),('mc_distrac','Distracciones'),
            ('mc_fij','1ª Fijación'),('mc_valencia','Valencia'),('mc_var','Var.Emoc.'),
        ]))
        pca.addWidget(self._eng_strip('engagement_bar2','_eng_lbl2',
                                      "Curva de engagement — segundo a segundo"))
        left_layout.addWidget(panel_cam, 1)

        cl.addWidget(left_widget, 55)

        # ── COLUMNA DERECHA: formulario ───────────────────────────
        right_widget = QWidget()
        right_widget.setStyleSheet(f"background-color:{BLANCO}; border-radius:8px; border:0.5px solid {GRIS_MEDIO};")
        rl = QVBoxLayout(right_widget); rl.setSpacing(0); rl.setContentsMargins(0,0,0,0)

        # Cabecera del formulario
        form_header = QWidget()
        form_header.setStyleSheet(f"background-color:{AZUL_OSCURO}; border-radius:8px 8px 0 0;")
        fhl = QHBoxLayout(form_header); fhl.setContentsMargins(10,8,10,8); fhl.setSpacing(10)

        # Foto del participante (se muestra al iniciar sesión)
        self.foto_header_lbl = QLabel()
        self.foto_header_lbl.setFixedSize(52, 52)
        self.foto_header_lbl.setStyleSheet(
            "border:2px solid rgba(255,255,255,0.3); border-radius:26px; "
            "background-color:rgba(255,255,255,0.1);")
        self.foto_header_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.foto_header_lbl.setText("👤")
        self.foto_header_lbl.setFont(QFont("Arial", 20))
        fhl.addWidget(self.foto_header_lbl)

        fht = QLabel("Datos del participante")
        fht.setFont(QFont("Arial",11,QFont.Weight.Bold))
        fht.setStyleSheet(f"color:{BLANCO};"); fhl.addWidget(fht)
        fhl.addStretch()
        rl.addWidget(form_header)

        # Cuerpo del formulario
        form_body = QWidget()
        fbl = QVBoxLayout(form_body); fbl.setSpacing(5); fbl.setContentsMargins(10,8,10,8)

        def sec_label(text):
            lbl = QLabel(text)
            lbl.setStyleSheet(
                f"font-size:7px; font-weight:bold; color:{AZUL_CLARO}; "
                f"background-color:#e8edf5; border-radius:2px; padding:2px 6px; "
                f"text-transform:uppercase; letter-spacing:1px;")
            return lbl

        def field_label(text):
            lbl = QLabel(text)
            lbl.setStyleSheet(f"font-size:9px; color:#555; font-weight:500; border:none;")
            return lbl

        FIELD_STYLE = (f"font-size:10px; padding:4px 6px; border:1px solid {GRIS_MEDIO}; "
                       f"border-radius:4px; background:#fafafa; color:{TEXTO}; height:26px;")
        RO_STYLE    = (f"font-size:10px; padding:4px 6px; border:1px solid {GRIS_MEDIO}; "
                       f"border-radius:4px; background:#e8edf5; color:{AZUL_CLARO}; "
                       f"font-weight:bold; height:26px;")
        HI_STYLE    = (f"font-size:10px; padding:4px 6px; border:1px solid #eda100; "
                       f"border-radius:4px; background:#fffbea; color:{TEXTO}; height:26px;")

        # ── Sección Identificación ────────────────────────────────
        fbl.addWidget(sec_label("Identificación"))

        row_id = QHBoxLayout(); row_id.setSpacing(6)
        col_id = QVBoxLayout(); col_id.setSpacing(2)
        col_id.addWidget(field_label("ID"))
        self.participant_input = QLineEdit(f"P{participant_counter[0]:03d}")
        self.participant_input.setReadOnly(True)
        self.participant_input.setStyleSheet(RO_STYLE)
        self.participant_input.setFixedWidth(65)
        col_id.addWidget(self.participant_input)
        row_id.addLayout(col_id)

        col_cod = QVBoxLayout(); col_cod.setSpacing(2)
        col_cod.addWidget(field_label("Código ECOTEC"))
        self.codigo_input = QLineEdit()
        self.codigo_input.setPlaceholderText("202001234")
        self.codigo_input.setStyleSheet(HI_STYLE)
        col_cod.addWidget(self.codigo_input)
        row_id.addLayout(col_cod)
        fbl.addLayout(row_id)

        fbl.addWidget(field_label("Nombre y apellido"))
        self.nombre_input = QLineEdit()
        self.nombre_input.setPlaceholderText("Nickole Lee Sagbay")
        self.nombre_input.setStyleSheet(FIELD_STYLE)
        fbl.addWidget(self.nombre_input)

        row_es = QHBoxLayout(); row_es.setSpacing(6)
        col_ed = QVBoxLayout(); col_ed.setSpacing(2)
        col_ed.addWidget(field_label("Edad"))
        self.edad_input = QLineEdit()
        self.edad_input.setPlaceholderText("22")
        self.edad_input.setMaxLength(2)
        self.edad_input.setStyleSheet(FIELD_STYLE)
        self.edad_input.setFixedWidth(44)
        col_ed.addWidget(self.edad_input)
        row_es.addLayout(col_ed)
        col_sx = QVBoxLayout(); col_sx.setSpacing(2)
        col_sx.addWidget(field_label("Sexo"))
        self.sexo_combo = QComboBox()
        self.sexo_combo.addItems(["— Selec. —","Masculino","Femenino","Prefiero no decir"])
        self.sexo_combo.setStyleSheet(FIELD_STYLE)
        col_sx.addWidget(self.sexo_combo)
        row_es.addLayout(col_sx)
        fbl.addLayout(row_es)

        # ── Sección Datos académicos ──────────────────────────────
        fbl.addWidget(sec_label("Datos académicos"))

        fbl.addWidget(field_label("Facultad"))
        self.facultad_combo = QComboBox()
        self.facultad_combo.addItem("— Seleccionar —")
        for fac in FACULTADES_CARRERAS:
            self.facultad_combo.addItem(fac)
        self.facultad_combo.setStyleSheet(FIELD_STYLE)
        self.facultad_combo.currentIndexChanged.connect(self._update_carreras)
        fbl.addWidget(self.facultad_combo)

        fbl.addWidget(field_label("Carrera"))
        self.carrera_combo = QComboBox()
        self.carrera_combo.addItem("— Seleccione facultad —")
        self.carrera_combo.setStyleSheet(FIELD_STYLE)
        fbl.addWidget(self.carrera_combo)

        fbl.addWidget(field_label("Semestre"))
        self.semestre_combo = QComboBox()
        self.semestre_combo.addItem("—")
        for s in SEMESTRES:
            self.semestre_combo.addItem(s)
        self.semestre_combo.setStyleSheet(FIELD_STYLE)
        fbl.addWidget(self.semestre_combo)

        # ── Sección Control técnico ───────────────────────────────
        fbl.addWidget(sec_label("Control técnico"))
        chk_row = QHBoxLayout(); chk_row.setSpacing(14)
        chk_row.addStretch()
        self.check_gafas = QCheckBox("Usa gafas")
        self.check_gafas.setStyleSheet(f"font-size:10px; color:{TEXTO};")
        chk_row.addWidget(self.check_gafas)
        self.check_lentes = QCheckBox("Lentes de contacto")
        self.check_lentes.setStyleSheet(f"font-size:10px; color:{TEXTO};")
        chk_row.addWidget(self.check_lentes)
        chk_row.addStretch()
        fbl.addLayout(chk_row)

        # ── Botones ───────────────────────────────────────────────
        for btn_attr, label, color, hover, slot in [
            ('btn_calibrate', '🎯  CALIBRAR',      AZUL_CLARO,  AZUL_MEDIO,   self.on_calibrate),
            ('btn_start',     '▶  INICIAR SESIÓN', VERDE,       VERDE_OSCURO, self.on_start),
            ('btn_stop',      '⏹  DETENER',        ROJO,        '#a93226',    self.on_stop),
        ]:
            btn = QPushButton(label); btn.setFixedHeight(30)
            btn.setStyleSheet(f"""QPushButton {{ background-color:{color}; color:{BLANCO};
                font-size:10px; font-weight:bold; border-radius:6px; }}
                QPushButton:hover {{ background-color:{hover}; }}
                QPushButton:disabled {{ background-color:{GRIS_MEDIO}; color:#888; }}""")
            btn.clicked.connect(slot); setattr(self, btn_attr, btn)
            fbl.addWidget(btn)
        self.btn_start.setEnabled(False); self.btn_stop.setEnabled(False)

        self.lbl_status = QLabel("Estado: Ingresa los datos y presiona CALIBRAR")
        self.lbl_status.setStyleSheet(f"color:gray; font-size:8px;"); fbl.addWidget(self.lbl_status)

        rl.addWidget(form_body)
        cl.addWidget(right_widget, 45)
        main_layout.addWidget(content, 1)

        # ── FOOTER ────────────────────────────────────────────────
        footer = QWidget(); footer.setFixedHeight(26)
        footer.setStyleSheet(f"background-color:{AZUL_OSCURO}; border-top:2px solid {VERDE};")
        fl = QHBoxLayout(footer); fl.setContentsMargins(14,0,14,0)
        fl.addWidget(QLabel("© 2026 Universidad ECOTEC — Ingeniería en Sistemas Inteligentes",
                    styleSheet=f"color:{GRIS_MEDIO}; font-size:9px;"))
        fl.addStretch()
        fl.addWidget(QLabel("Campus Samborondón",
                    styleSheet=f"color:{GRIS_MEDIO}; font-size:9px;"))
        main_layout.addWidget(footer)

    def _update_carreras(self, index):
        """Actualiza el combo de carrera según la facultad seleccionada."""
        self.carrera_combo.clear()
        if index <= 0:
            self.carrera_combo.addItem("— Seleccione facultad —")
            return
        facultad = self.facultad_combo.currentText()
        carreras  = FACULTADES_CARRERAS.get(facultad, [])
        self.carrera_combo.addItem("— Seleccionar —")
        for c in carreras:
            self.carrera_combo.addItem(c)

    def on_camera_frame(self, frame, data):
        if frame is None:
            return
        if self.session_active and data and data['is_looking'] and self.video_player_obj:
            raw_x, raw_y = self.video_player_obj.gaze_x, self.video_player_obj.gaze_y
            rw, rh = self._heatmap_rw, self._heatmap_rh
            if rw > 0 and rh > 0:
                cx = int(np.clip(raw_x/rw*640, 0, 639))
                cy = int(np.clip(raw_y/rh*400, 0, 399))
            else:
                cx, cy = 320, 200
            cv2.circle(self.heatmap, (cx,cy), 30, 1.0, -1)
            self.heatmap = cv2.GaussianBlur(self.heatmap, (51,51), 0)
            self.heatmap = np.clip(self.heatmap, 0, 1)

        # Actualizar panel cámara
        rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
        h, w, ch = rgb.shape
        qt_image  = QImage(rgb.data, w, h, ch*w, QImage.Format.Format_RGB888)
        lw = self.camera_label.width(); lh = self.camera_label.height()
        if lw > 0 and lh > 0:
            self.camera_label.setPixmap(QPixmap.fromImage(
                qt_image.scaled(lw, lh, Qt.AspectRatioMode.KeepAspectRatio,
                                Qt.TransformationMode.FastTransformation)))

        if data:
            if self.session_active and self.video_player_obj:
                self.video_player_obj.set_emotion(data['emotion'])
                self.video_player_obj.update_gaze(data['yaw'], data['pitch'])

            emotion_icons = {'happy':'😊','sad':'😢','angry':'😠',
                             'fear':'😨','surprise':'😲','disgust':'🤢','neutral':'😐'}
            icon = emotion_icons.get(data['emotion'], '❓')

            # Barra izquierda (panel comercial)
            self.mc_emocion.setText(f"{icon}")
            self.mc_yaw.setText(f"{data['yaw']:.1f}°")
            self.mc_pitch.setText(f"{data['pitch']:.1f}°")
            col = VERDE if data['is_looking'] else ROJO
            self.mc_pantalla.setText("✓ SÍ" if data['is_looking'] else "✗ NO")
            self.mc_pantalla.setStyleSheet(f"color:{col}; border:none;")
            fps_val = data.get('fps', 0.0)
            fps_col = VERDE if fps_val >= 20 else ROJO
            self.mc_fps.setText(f"{fps_val:.0f}")
            self.mc_fps.setStyleSheet(f"color:{fps_col}; border:none; font-weight:bold;")

            if self.session_active:
                self.engagement_bar.setValue(int(data['engagement']*100))

            if self.session_active and self.video_player_obj:
                sector = self.video_player_obj.current_sector
                self.mc_sector.setText(str(sector))
                if self.fullscreen_video:
                    self.fullscreen_video.update_debug_info(
                        sector, data['yaw'], data['pitch'],
                        data['emotion'], data.get('fps', 0.0))

    def on_calibrate(self):
        nombre = self.nombre_input.text().strip()
        if not nombre:
            self.lbl_status.setText("Estado: Ingresa el nombre antes de calibrar"); return
        dist_dialog = DistanceDialog(self.controller, self)
        if dist_dialog.exec() != QDialog.DialogCode.Accepted: return
        params = dist_dialog.get_geometry_params()
        self.controller.pitch_offset    = params['pitch_offset']
        self.controller.yaw_threshold   = params['yaw_threshold']
        self.controller.pitch_threshold = params['pitch_threshold']
        self._participant_distancia     = params['distance_cm']
        self.mc_dist.setText(f"{params['distance_cm']}cm")
        self.lbl_status.setText(
            f"Estado: Calibrando a {params['distance_cm']}cm — mire cada punto rojo...")
        from utils.calibration import CalibrationWidget
        screen_idx = 1 if len(QApplication.screens()) > 1 else 0
        self.calibration_widget = CalibrationWidget(self.controller, screen_index=screen_idx)
        self.calibration_widget.calibration_finished.connect(self.on_calibration_done)
        self.btn_calibrate.setEnabled(False)

    def on_calibration_done(self, matrix):
        self.controller.calibration_matrix = matrix
        self.btn_start.setEnabled(True); self.btn_calibrate.setEnabled(True)
        self.lbl_status.setText("Estado: Calibración completada ✅ — listo para iniciar sesión")
        self.raise_(); self.activateWindow()

    def on_start(self):
        nombre = self.nombre_input.text().strip()
        if not nombre:
            self.lbl_status.setText("Estado: Ingresa el nombre del participante"); return

        # Validar carrera seleccionada
        carrera_sel = self.carrera_combo.currentText()
        if carrera_sel in ('— Seleccionar —', '— Seleccione facultad —', ''):
            self.lbl_status.setText("Estado: Selecciona la carrera del participante")
            QMessageBox.warning(self, "Campo requerido",
                "Por favor selecciona la carrera del participante antes de iniciar.")
            return

        # Validar código ECOTEC
        codigo_sel = self.codigo_input.text().strip()
        if not codigo_sel:
            self.lbl_status.setText("Estado: Ingresa el código ECOTEC del participante")
            QMessageBox.warning(self, "Campo requerido",
                "Por favor ingresa el código ECOTEC del participante antes de iniciar.")
            return

        participant_id = self.participant_input.text()
        self.controller.start_session(participant_id)

        self.btn_start.setEnabled(False)
        self.btn_calibrate.setEnabled(False)
        self.btn_stop.setEnabled(True)
        self.session_active = True

        # Recoger todos los campos del formulario
        self._participant_name     = nombre
        self._participant_codigo   = self.codigo_input.text().strip()
        self._participant_edad     = self.edad_input.text().strip()
        self._participant_sexo     = self.sexo_combo.currentText()
        self._participant_facultad = self.facultad_combo.currentText()
        self._participant_carrera  = self.carrera_combo.currentText()
        self._participant_semestre = self.semestre_combo.currentText()
        self._participant_gafas    = self.check_gafas.isChecked()
        self._participant_lentes   = self.check_lentes.isChecked()

        self._participant_photo = self._capture_photo(participant_id, nombre)
        self.heatmap = np.zeros((400, 640), dtype=np.float32)
        self.lbl_id_header.setText(f"● Sesión activa: {nombre} — {participant_id}")

        # Mostrar foto en el header del formulario
        if self._participant_photo and os.path.exists(self._participant_photo):
            try:
                pix = QPixmap(self._participant_photo).scaled(
                    52, 52,
                    Qt.AspectRatioMode.KeepAspectRatioByExpanding,
                    Qt.TransformationMode.SmoothTransformation)
                self.foto_header_lbl.setPixmap(pix)
                self.foto_header_lbl.setText("")
            except Exception:
                pass
        self.lbl_status.setText(f"Estado: Grabando — {nombre} ({participant_id})")

        import glob
        videos = (glob.glob('videos/*.mp4') + glob.glob('videos/*.webm') + glob.glob('videos/*.mkv'))
        if not videos:
            self.lbl_status.setText("Estado: No se encontró video en carpeta videos/"); return
        video_path = videos[0]
        audio_path = os.path.splitext(video_path)[0] + '.mp3'

        self.fullscreen_video = FullscreenVideoWindow(screen_index=self._target_screen_index())
        from utils.video_player import VideoPlayer
        self.video_player_obj = VideoPlayer(
            video_path,
            audio_path=audio_path if os.path.exists(audio_path) else None)
        self.video_player_obj.video_finished.connect(self.on_stop)
        self.video_player_obj.timestamp_updated.connect(self.on_video_timestamp)
        self.video_player_obj.video_thread_ready.connect(self._connect_video_frames)
        self.video_player_obj.load()
        if hasattr(self.controller,'calibration_matrix') and self.controller.calibration_matrix:
            self.video_player_obj.set_calibration(self.controller.calibration_matrix)
        QTimer.singleShot(100, self._configure_render_size)
        QTimer.singleShot(400, self.video_player_obj.play)
        if not os.path.exists(audio_path):
            self.audio_thread = AudioThread(video_path); self.audio_thread.start()
        QTimer.singleShot(500, self.raise_)
        QTimer.singleShot(500, self.activateWindow)

    def _configure_render_size(self):
        if self.fullscreen_video:
            rw, rh = self.fullscreen_video.get_render_size()
            if rw <= 0 or rh <= 0:
                geom = self.fullscreen_video.geometry()
                rw, rh = geom.width(), geom.height()
        else:
            rw = self.comercial_label.width()
            rh = self.comercial_label.height()
        if rw > 0 and rh > 0:
            self.video_player_obj.set_render_size(rw, rh)
            self._heatmap_rw = rw; self._heatmap_rh = rh
            print(f"Render size: {rw}x{rh}")

    def _connect_video_frames(self, thread):
        thread.frame_ready.connect(self.on_video_frame)

    def on_video_frame(self, frame, current_time):
        rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
        h, w, ch = rgb.shape
        qt_image  = QImage(rgb.data, w, h, ch*w, QImage.Format.Format_RGB888)
        lw = self.comercial_label.width(); lh = self.comercial_label.height()
        if lw > 0 and lh > 0:
            self.comercial_label.setPixmap(QPixmap.fromImage(
                qt_image.scaled(lw, lh, Qt.AspectRatioMode.KeepAspectRatio,
                                Qt.TransformationMode.FastTransformation)))
        if self.fullscreen_video:
            self.fullscreen_video.update_frame(qt_image)
        mins = int(current_time // 60); secs = int(current_time % 60)
        # Actualizar badge del panel comercial
        try:
            # Actualizar el texto del header dinámicamente no es trivial;
            # solo mostramos el tiempo en el label de status
            self.lbl_status.setText(
                f"Estado: Grabando — {self._participant_name} — {mins:02d}:{secs:02d}")
        except Exception:
            pass

    def on_video_timestamp(self, time, sector, emotion):
        pass

    def on_stop(self):
        """Detiene la sesión y guarda automáticamente todos los archivos."""
        self.session_active = False
        sector_log = []
        aoi_log    = []
        if self.video_player_obj:
            sector_log = self.video_player_obj.sector_log
            aoi_log    = self.video_player_obj.aoi_log
            self.video_player_obj.stop(); self.video_player_obj = None
        if self.fullscreen_video:
            self.fullscreen_video.close_window(); self.fullscreen_video = None
        self.comercial_label.clear()
        self.comercial_label.setText("▶  Calibra y luego inicia sesión para reproducir el comercial")
        self.comercial_label.setStyleSheet(f"background-color:#111; color:{GRIS_MEDIO}; border:none;")

        filename, summary = self.controller.stop_session()
        self.btn_start.setEnabled(True)
        self.btn_calibrate.setEnabled(True)
        self.btn_stop.setEnabled(False)
        self.lbl_status.setText("Estado: Guardando archivos...")

        # ── GUARDADO AUTOMÁTICO ────────────────────────────────────
        participant_id   = self.participant_input.text()
        nombre           = self._participant_name
        nombre_safe      = _nombre_safe(nombre)
        base             = f"data/{nombre_safe}_{participant_id}"

        participant_data = {
            'nombre':     nombre,
            'codigo':     self._participant_codigo,
            'edad':       self._participant_edad,
            'sexo':       self._participant_sexo,
            'facultad':   self._participant_facultad,
            'carrera':    self._participant_carrera,
            'semestre':   self._participant_semestre,
            'gafas':      self._participant_gafas,
            'lentes':     self._participant_lentes,
            'distancia':  self._participant_distancia,
            'photo_path': self._participant_photo,
            'fecha':      summary.get('fecha', datetime.now().strftime('%d/%m/%Y')),
            'hora':       summary.get('hora',  datetime.now().strftime('%H:%M')),
        }

        saved_files = {}
        os.makedirs('data', exist_ok=True)
        try:
            from utils.export_utils import (save_individual_excel,
                                      save_individual_pdf,
                                      save_consolidated_excel)
            # 1. Excel individual
            xl_path = f"{base}.xlsx"
            save_individual_excel(
                participant_id, participant_data, summary, sector_log,
                self.heatmap.copy(), self._participant_photo, xl_path,
                aoi_log=aoi_log)
            saved_files['Excel individual'] = xl_path

            # 2. PDF individual
            pdf_path = f"{base}.pdf"
            save_individual_pdf(
                participant_id, participant_data, summary, sector_log,
                self.heatmap.copy(), self._participant_photo, pdf_path,
                aoi_log=aoi_log)
            saved_files['PDF individual'] = pdf_path

            # 3. Excel consolidado
            cons_path = 'data/reporte_consolidado.xlsx'
            save_consolidated_excel(
                participant_id, participant_data, summary, sector_log, cons_path,
                aoi_log=aoi_log)
            saved_files['Excel consolidado'] = cons_path

            self.lbl_status.setText(f"Estado: ✅ Archivos guardados — {nombre}")
        except Exception as e:
            print(f"Error guardando archivos: {e}")
            self.lbl_status.setText(f"Estado: ⚠ Error al guardar — {e}")

        # Avanzar contador y limpiar formulario
        participant_counter[0] += 1
        self.participant_input.setText(f"P{participant_counter[0]:03d}")
        self.nombre_input.clear()
        self.codigo_input.clear()
        self.edad_input.clear()
        self.sexo_combo.setCurrentIndex(0)
        self.facultad_combo.setCurrentIndex(0)
        self.carrera_combo.clear()
        self.carrera_combo.addItem("— Seleccione facultad —")
        self.semestre_combo.setCurrentIndex(0)
        self.check_gafas.setChecked(False)
        self.check_lentes.setChecked(False)
        self.lbl_id_header.setText("")
        self.foto_header_lbl.setPixmap(QPixmap())
        self.foto_header_lbl.setText("👤")
        self.raise_(); self.activateWindow()

        # Mostrar diálogo de resultados
        dialog = ResultsDialog(
            summary, self.heatmap, sector_log,
            participant_name=nombre,
            saved_files=saved_files,
            parent=self,
            aoi_log=aoi_log)
        dialog.exec()


# ─────────────────────────────────────────────
#  PANTALLA DE CARGA
# ─────────────────────────────────────────────
class LoadingScreen(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Cargando...")
        self.setFixedSize(500, 240)
        self.setWindowFlags(Qt.WindowType.FramelessWindowHint)
        self._build_ui(); self._start_loading()

    def _build_ui(self):
        central = QWidget(); self.setCentralWidget(central)
        central.setStyleSheet(f"background-color:{AZUL_OSCURO};")
        layout = QVBoxLayout(central)
        layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.setSpacing(14); layout.setContentsMargins(30,30,30,30)
        for text, size, color in [
            ("Sistema de Eye Tracking", 20, BLANCO),
            ("Análisis de Engagement Publicitario", 11, GRIS_MEDIO),
            ("Universidad ECOTEC — Ingeniería en Sistemas Inteligentes", 9, VERDE),
        ]:
            lbl = QLabel(text)
            lbl.setFont(QFont("Arial", size,
                              QFont.Weight.Bold if size==20 else QFont.Weight.Normal))
            lbl.setStyleSheet(f"color:{color};")
            lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
            layout.addWidget(lbl)
        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0,0); self.progress_bar.setFixedHeight(6)
        self.progress_bar.setStyleSheet(f"""
            QProgressBar {{ background-color:#001040; border-radius:3px; border:none; }}
            QProgressBar::chunk {{ background-color:{VERDE}; border-radius:3px; }}""")
        layout.addWidget(self.progress_bar)
        self.status_label = QLabel("Iniciando sistema...")
        self.status_label.setStyleSheet(f"color:{GRIS_MEDIO}; font-size:10px;")
        self.status_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(self.status_label)

    def _start_loading(self):
        QTimer.singleShot(300, self._begin_thread)

    def _begin_thread(self):
        self.thread = LoadingThread()
        self.thread.progress.connect(self.status_label.setText)
        self.thread.finished.connect(self.on_loaded)
        self.thread.start()

    def on_loaded(self, controller):
        self.main_window = MainView(controller)
        self.main_window.show()
        self.main_window.start_camera_thread()
        self.close()


def main():
    app = QApplication(sys.argv)
    loading = LoadingScreen()
    loading.show()
    sys.exit(app.exec())


if __name__ == '__main__':
    main()